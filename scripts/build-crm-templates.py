#!/usr/bin/env python3
"""
build-crm-templates.py —— 把用户真实报价单/合同 docx 注入 docxtemplater 标签，
生成 resources/crm-templates/{quotation,contract}.docx（提交进仓库）。

用法：python3 scripts/build-crm-templates.py [源目录]
默认源目录 /tmp/crm-tpl-inspect（真实模版副本）。重复运行覆盖输出。

注入规则（与 crmDocGenCore.buildDocData 的数据字段一一对应）：
- quotation.docx table0:
    r1 数据行 → 循环行 {#items}{idx}|{name}|{spec}|{qty}|{unit_price}|{subtotal}|{remark}{/items}
    删 r2 空行；r3 合计 金额格→{total} 备注格→{footer_remark}；r4 含运费行保留
- contract.docx:
    p2 甲方（买方）：{buyer_name} … 合同编号：{no}；p4 签定时间：{sign_date}
    table0 r1 → 循环行 {#items}{name}|{spec}|{unit}|{qty}|{unit_price}|{amount}|{remark}{/items}
    删 r2 空行；r3 合计 金额格→{total}；r4 大写行→合计人民币金额（大写）：{amount_cn}
    table1 甲方块：{buyer_name}/{buyer_addr}/{buyer_bank}/{buyer_account}/{buyer_tax}/{buyer_phone}
"""
import os
import sys
import docx


def set_cell_text(cell, text: str) -> None:
    """第一个段落写入 text（单 run），移除多余段落。"""
    cell.paragraphs[0].text = text
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def delete_row(table, idx: int) -> None:
    tr = table.rows[idx]._tr
    tr.getparent().remove(tr)


def build_quotation(src: str, dst: str) -> None:
    d = docx.Document(src)
    # To 客户行：run0 客户名 → {customer}，保留换行与「报价目录/Quotation」子标题
    d.paragraphs[3].runs[0].text = 'To：{customer}  '
    t = d.tables[0]
    # r1 数据行 → 循环行
    cells = t.rows[1].cells
    set_cell_text(cells[0], '{#items}{idx}')
    set_cell_text(cells[1], '{name}')
    set_cell_text(cells[2], '{spec}')
    set_cell_text(cells[3], '{qty}')
    set_cell_text(cells[4], '{unit_price}')
    set_cell_text(cells[5], '{subtotal}')
    set_cell_text(cells[7], '{remark}{/items}')
    # r2 空行删除
    delete_row(t, 2)
    # r3 合计行：金额格→{total}，备注格→{footer_remark}
    cells = t.rows[2].cells
    set_cell_text(cells[5], '{total}')
    set_cell_text(cells[7], '{footer_remark}')
    d.save(dst)
    print(f'[ok] quotation → {dst}')


def build_contract(src: str, dst: str) -> None:
    d = docx.Document(src)
    # p2/p4
    d.paragraphs[2].text = '甲方（买方）：{buyer_name}            合同编号：{no}'
    d.paragraphs[4].text = '乙方（卖方）：无锡库叉搬运设备有限公司               签定时间：{sign_date}'
    # table0 数据行 → 循环行
    t = d.tables[0]
    cells = t.rows[1].cells
    set_cell_text(cells[0], '{#items}{name}')
    set_cell_text(cells[1], '{spec}')
    set_cell_text(cells[2], '{unit}')
    set_cell_text(cells[3], '{qty}')
    set_cell_text(cells[4], '{unit_price}')
    set_cell_text(cells[5], '{amount}')
    set_cell_text(cells[6], '{remark}{/items}')
    # r2 空行删除
    delete_row(t, 2)
    # r3 合计 金额格→{total}
    set_cell_text(t.rows[2].cells[5], '{total}')
    # r4 大写行
    set_cell_text(t.rows[3].cells[0], '合计人民币金额（大写）：{amount_cn}')
    # table1 甲方块
    t1 = d.tables[1]
    acells = t1.rows[0].cells[0].paragraphs
    acells[0].text = '甲方：{buyer_name}'
    acells[1].text = '单位地址:{buyer_addr}'
    acells[2].text = '开户银行：{buyer_bank}'
    acells[3].text = '账  号：{buyer_account}'
    acells[4].text = '税  号：{buyer_tax}'
    acells[5].text = '电  话：{buyer_phone}'
    # acells[6] 授权代表签署：保留
    d.save(dst)
    print(f'[ok] contract → {dst}')


def main() -> None:
    src_dir = sys.argv[1] if len(sys.argv) > 1 else '/tmp/crm-tpl-inspect'
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out_dir = os.path.join(project_root, 'resources', 'crm-templates')
    os.makedirs(out_dir, exist_ok=True)
    build_quotation(os.path.join(src_dir, 'quotation.docx'),
                    os.path.join(out_dir, 'quotation.docx'))
    build_contract(os.path.join(src_dir, 'contract.docx'),
                   os.path.join(out_dir, 'contract.docx'))


if __name__ == '__main__':
    main()
