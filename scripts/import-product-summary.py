#!/usr/bin/env python3
"""
import-product-summary.py —— 产品汇总表 → CRM 产品库一次性导入
用法: python3 scripts/import-product-summary.py <xlsx路径> [--db <weflow-crm.db路径>]
来源表结构：产品汇总对比 / 库叉详细技术参数 / 库叉产品性能特点 / 外调报价原始明细
定价口径（用户确认 2026-08-18）：库叉自产=含税运；外调车型=裸车价不含税运。
X1-Li 与 X1C-Li 为同一产品（按 X1C-Li 导入）。
"""
import json, re, sqlite3, sys, time
from pathlib import Path

import openpyxl

PRICE_CORRECTIONS = {'定制款': '5500-6000'}  # 源表笔误 5500-600
INQUIRE_KEYWORDS = ('详询', '按车型报价')

def parse_price(raw, model):
    """返回 (unit_price, variants, note)。variants=[{label, price, price_max?}]"""
    text = PRICE_CORRECTIONS.get(model) or str(raw or '').strip()
    text = text.replace(',', '').replace('，', '/')
    if not text or text in ('—', '--') or any(k in text for k in INQUIRE_KEYWORDS):
        return None, [], ('按车型报价（详询）' if any(k in text for k in INQUIRE_KEYWORDS) else '')
    variants = []
    # 多规格：2T:11000 / 3T:12000 或 20AH:9000 / 32AH:9700
    parts = [p.strip() for p in re.split(r'[/／]', text) if p.strip()]
    labeled = [p for p in parts if re.match(r'^[\d.]+(?:T|AH)\s*[:：]', p, re.I)]
    if labeled and len(labeled) == len(parts) and len(parts) > 1:
        for p in parts:
            m = re.match(r'^([\d.]+(?:T|AH))\s*[:：]\s*([\d.]+)(?:\s*[-~]\s*([\d.]+))?', p, re.I)
            if not m:
                continue
            v = {'label': m.group(1).upper(), 'price': float(m.group(2))}
            if m.group(3): v['price_max'] = float(m.group(3))
            variants.append(v)
        if variants:
            return min(v['price'] for v in variants), variants, ''
    # 单值或区间：3400 / 5500-6000
    m = re.match(r'^([\d.]+)(?:\s*[-~]\s*([\d.]+))?$', text)
    if m:
        price = float(m.group(1))
        if m.group(2):
            return price, [{'label': '区间', 'price': price, 'price_max': float(m.group(2))}], ''
        return price, [], ''
    return None, [], f'价格待解析：{text}'

def derive_subcategory(name):
    if '前移' in name: return '前移式堆高车'
    if '座驾' in name: return '座驾式堆高车'
    if '步行配重' in name: return '步行配重式堆高车'
    if '配重' in name: return '配重式堆高车'
    if '手动' in name and '不锈钢' in name: return '不锈钢手动搬运车'
    if '手动' in name: return '手动搬运车'
    if '电子秤' in name: return '电子秤搬运车'
    if '半电动' in name: return '半电动堆高车'
    if '步行式堆高' in name: return '步行式堆高车'
    if '堆高' in name: return '全电动堆高车'
    if '搬运' in name: return '全电动搬运车'
    return ''

def main():
    if len(sys.argv) < 2:
        sys.exit('用法: python3 import-product-summary.py <xlsx路径> [--db <db路径>]')
    xlsx_path = Path(sys.argv[1])
    db_path = Path(sys.argv[3]) if len(sys.argv) > 3 and sys.argv[2] == '--db' else \
        Path.home() / 'Library/Application Support/WeFlow/weflow-crm.db'
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # ── 详细技术参数（自产 6 车型 × N 参数）──
    detail = {}
    ws2 = wb['库叉详细技术参数']
    rows2 = list(ws2.iter_rows(values_only=True))
    if rows2:
        header = [str(c or '') for c in rows2[0]]
        for r in rows2[1:]:
            item, unit = str(r[0] or '').strip(), str(r[1] or '').strip()
            if not item: continue
            key = f'{item}({unit})' if unit and unit != '—' else item
            for idx, col in enumerate(header[2:], start=2):
                if idx < len(r) and r[idx] not in (None, '', '—'):
                    detail.setdefault(col, {})[key] = str(r[idx]).strip()

    # ── 性能特点 ──
    features = {}
    ws3 = wb['库叉产品性能特点']
    for r in ws3.iter_rows(values_only=True):
        if r and r[0] and str(r[0]) != '产品型号' and r[1]:
            features[str(r[0]).strip()] = str(r[1]).strip()

    def match_detail(model):
        for col in detail:
            if col.startswith(model.split('(')[0]): return detail[col]
        return None
    def match_feature(model):
        for k, v in features.items():
            if k.startswith(model.split('(')[0]): return v
        return None

    # ── 主表 ──
    ws1 = wb['产品汇总对比']
    cols = [str(c.value or '') for c in ws1[1]]
    products = []
    for row in ws1.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]: continue
        src = str(row[0] or '').strip()
        model = str(row[1] or '').strip()
        name = str(row[2] or model).strip()
        price_term = '含税运' if src == '库叉自产' else '裸车价（不含税运）'
        unit_price, variants, price_note = parse_price(row[16], model)
        summary = {cols[i]: str(row[i]).strip() for i in range(len(cols)) if i < len(row) and row[i] not in (None, '')}
        specs = {'参数': summary, '价格口径': price_term}
        d = match_detail(model)
        if d: specs['详细技术参数'] = d
        if price_note: specs['价格备注'] = price_note
        desc_parts = []
        feat = match_feature(model)
        if feat: desc_parts.append(feat)
        if row[17]: desc_parts.append(str(row[17]).strip())
        desc_parts.append(f'报价口径：{price_term}')
        products.append({
            'model': model, 'name': name,
            'spec': f"载荷{row[5] or '—'}｜{row[3] or '—'}｜{row[4] or ''}{(' ' + str(row[7])) if row[7] else ''}｜举升{row[9] or '—'}",
            'unit_price': unit_price, 'reference_price': unit_price,
            'category': src, 'subcategory': derive_subcategory(name),
            'product_line': src,
            'description': '\n'.join(desc_parts),
            'specs': json.dumps(specs, ensure_ascii=False),
            'variants': json.dumps(variants, ensure_ascii=False),
        })

    # ── 写库（整表替换：旧 4 条测试数据全部被新表覆盖）──
    con = sqlite3.connect(db_path)
    cur = con.cursor()
    before = cur.execute('SELECT COUNT(*) FROM product').fetchone()[0]
    cur.execute('DELETE FROM product')
    now = int(time.time() * 1000)
    for p in products:
        cur.execute(
            'INSERT INTO product (model, name, spec, unit_price, product_line, created_at, category, subcategory, reference_price, description, specs, variants) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
            (p['model'], p['name'], p['spec'], p['unit_price'], p['product_line'], now, p['category'], p['subcategory'], p['reference_price'], p['description'], p['specs'], p['variants']))
    con.commit()
    after = cur.execute('SELECT COUNT(*) FROM product').fetchone()[0]
    no_price = [p['model'] for p in products if p['unit_price'] is None]
    print(f'导入完成：{before} → {after} 个产品（自产 {sum(1 for p in products if p["category"]=="库叉自产")} / 外调 {sum(1 for p in products if p["category"]=="外调车型")}）')
    if no_price: print(f'无固定价格（详询）: {", ".join(no_price)}')
    for p in products:
        v = json.loads(p['variants'])
        price_txt = f"¥{p['unit_price']:.0f}" if p['unit_price'] is not None else '详询'
        print(f"  {p['category'][:2]} | {p['model'][:24]:<24} | {price_txt:<9} | {p['subcategory']}" + (f' | variants:{len(v)}' if v else ''))

if __name__ == '__main__':
    main()
